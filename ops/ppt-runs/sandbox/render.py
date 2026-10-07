"""Collect a run's deliverables and render them to PNG pages, with objective
editability counts for .pptx files. Usage: render.py <workdir> <outdir>."""
import json, shutil, subprocess, sys
from pathlib import Path

DELIVERABLE = {".pptx", ".html", ".pdf", ".png", ".jpg", ".jpeg", ".key", ".md"}
MAX_PAGES = 12


def walk(shapes):
    """Every shape, including those inside groups."""
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:   # group
            yield from walk(sh.shapes)


def pptx_stats(path: Path) -> dict:
    from pptx import Presentation
    prs = Presentation(str(path)); area = prs.slide_width * prs.slide_height
    slides, text_chars, text_shapes, pictures, full_image_slides = len(prs.slides), 0, 0, 0, 0
    for s in prs.slides:
        for sh in walk(s.shapes):
            if sh.has_text_frame and sh.text_frame.text.strip():
                text_chars += len(sh.text_frame.text.strip()); text_shapes += 1
            if getattr(sh, "has_table", False) and sh.has_table:
                text_chars += sum(len(c.text.strip()) for r in sh.table.rows for c in r.cells); text_shapes += 1
            if sh.shape_type == 13:
                pictures += 1
                if sh.width * sh.height > 0.8 * area: full_image_slides += 1
    return {"slides": slides, "text_chars": text_chars, "text_shapes": text_shapes,
            "pictures": pictures, "full_image_slides": full_image_slides}


SLIDE_SELECTORS = ["section.slide", ".slide", "section", "[data-slide]", ".reveal .slides > section"]
VIEWPORT = {"width": 1600, "height": 900}
SETTLE_MS = 1500        # entrance animations finish before the screenshot


def html_pages(src: Path, pages: Path) -> int:
    """Screenshot a web deck slide by slide in a real browser. Printing to PDF loses
    content that appears only through scroll or entrance animations. Decks turn pages
    either on a keypress (an "active" slide) or by scrolling: press ArrowRight first,
    and when the screen does not change, scroll the next slide into view."""
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path="/usr/bin/chromium", args=["--no-sandbox"])
        page = b.new_page(viewport=VIEWPORT); page.goto(src.as_uri()); page.wait_for_timeout(SETTLE_MS)
        slides = next((page.query_selector_all(s) for s in SLIDE_SELECTORS if len(page.query_selector_all(s)) > 1), [])
        n, last = min(len(slides) or MAX_PAGES, MAX_PAGES), None
        for i in range(n):
            if i:
                page.keyboard.press("ArrowRight"); page.wait_for_timeout(SETTLE_MS)
                if page.screenshot() == last and slides:
                    slides[i].scroll_into_view_if_needed(); page.wait_for_timeout(SETTLE_MS)
            last = page.screenshot(path=str(pages / f"p-{i + 1:02d}.png"))
        b.close()
    return n


def to_pdf(src: Path, outdir: Path) -> Path | None:
    if src.suffix == ".pptx":
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(src)], capture_output=True, timeout=300)
        return outdir / (src.stem + ".pdf")
    return src if src.suffix == ".pdf" else None


def main() -> None:
    work, out = Path(sys.argv[1]), Path(sys.argv[2])
    files = [p for p in work.rglob("*") if p.is_file() and p.suffix.lower() in DELIVERABLE and p.name != "brief.md" and "node_modules" not in p.parts]
    deliver = out / "deliverables"; deliver.mkdir(exist_ok=True); pages = out / "pages"; pages.mkdir(exist_ok=True)
    report = {"files": [], "pptx": {}}
    for f in files:
        shutil.copy2(f, deliver / f.name); report["files"].append(str(f.relative_to(work)))
        if f.suffix == ".pptx":
            try: report["pptx"][f.name] = pptx_stats(f)
            except Exception as e: report["pptx"][f.name] = {"error": str(e)[:200]}
    main_file = next((f for ext in (".pptx", ".html", ".pdf") for f in files if f.suffix == ext), None)
    if main_file and main_file.suffix == ".html":
        html_pages(main_file, pages); pdf = None
    else:
        pdf = to_pdf(main_file, out) if main_file else None
    if pdf and pdf.exists():
        subprocess.run(["pdftoppm", "-png", "-r", "60", "-l", str(MAX_PAGES), str(pdf), str(pages / "p")], capture_output=True, timeout=300)
    report["rendered_from"] = main_file.name if main_file else None
    report["pages"] = sorted(p.name for p in pages.glob("*.png"))
    (out / "render.json").write_text(json.dumps(report, indent=1))


if __name__ == "__main__":
    main()
